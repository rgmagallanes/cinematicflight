from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "cinematic-flight-sales-playbook.docx"
LOGO = ROOT / "public" / "assets" / "brand" / "cinematic-flight-logo-concept-v1.png"
HERO = ROOT / "public" / "assets" / "property" / "good-food-farm-overview-1778.png"

FOREST = RGBColor(0x07, 0x11, 0x0D)
FOREST_SOFT = RGBColor(0x10, 0x20, 0x19)
PAPER = RGBColor(0xEE, 0xE9, 0xDF)
PAPER_DEEP = RGBColor(0xDE, 0xD6, 0xC7)
INK = RGBColor(0x14, 0x20, 0x19)
BRASS = RGBColor(0xC7, 0x97, 0x49)
BRASS_DARK = RGBColor(0x8B, 0x63, 0x2D)
SAGE = RGBColor(0x91, 0xA0, 0x96)
WHITE = RGBColor(0xF5, 0xF0, 0xE5)


def set_cell_fill(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color_hex)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="CFC6B5", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)
        borders.append(el)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_run(run, font="Avenir Next", size=10.5, color=INK, bold=False, italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    return run


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_kicker(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text.upper()), size=8.5, color=BRASS_DARK, bold=True)
    p.runs[0].font.letter_spacing = Pt(1.1)
    return p


def add_h1(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    keep_with_next(p)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(text, style="Heading 2")
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None, after=8, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.24
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]), italic=italic)
    else:
        set_run(p.add_run(text), italic=italic)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))
    return p


def add_number(doc, title, detail, style="List Number"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    set_run(p.add_run(title + " — "), bold=True, color=BRASS_DARK)
    set_run(p.add_run(detail))
    return p


def add_quote(doc, text, label=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(6.42)
    cell = table.cell(0, 0)
    cell.width = Inches(6.42)
    set_cell_fill(cell, "E7DDC9")
    set_cell_margins(cell, top=170, start=220, bottom=170, end=220)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0 if label is None else 5)
    p.paragraph_format.line_spacing = 1.18
    set_run(p.add_run(text), font="Bodoni Moda", size=15.5, color=FOREST, italic=False)
    if label:
        lp = cell.add_paragraph()
        lp.paragraph_format.space_after = Pt(0)
        set_run(lp.add_run(label.upper()), size=7.8, color=BRASS_DARK, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_note(doc, label, text, fill="102019", text_color=WHITE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.42)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, top=150, start=190, bottom=150, end=190)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(label.upper()), size=8, color=BRASS, bold=True)
    q = cell.add_paragraph()
    q.paragraph_format.space_after = Pt(0)
    q.paragraph_format.line_spacing = 1.18
    set_run(q.add_run(text), size=10.2, color=text_color)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_editable(doc, label, prompt):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run(label + ": "), size=9.5, color=BRASS_DARK, bold=True)
    run = set_run(p.add_run(f"[{prompt}]"), size=9.5, color=FOREST, italic=True)
    run.font.highlight_color = None
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3E9D5")
    p_pr.append(shd)
    return p


def page_break(doc):
    doc.add_page_break()


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.92)
    section.right_margin = Inches(0.92)
    section.header_distance = Inches(0.34)
    section.footer_distance = Inches(0.34)

    normal = doc.styles["Normal"]
    normal.font.name = "Avenir Next"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Avenir Next")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Avenir Next")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.24

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Bodoni Moda"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Bodoni Moda")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Bodoni Moda")
    h1.font.size = Pt(25)
    h1.font.bold = False
    h1.font.color.rgb = FOREST
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Avenir Next"
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Avenir Next")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Avenir Next")
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = BRASS_DARK
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Avenir Next"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Avenir Next")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Avenir Next")
        style.font.size = Pt(10.5)
        style.font.color.rgb = INK

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hp.add_run("CINEMATIC FLIGHT  /  SALES PLAYBOOK"), size=7.6, color=SAGE, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp.add_run("Living document  •  Update the brass fields as the offer develops  •  "), size=7.3, color=SAGE)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)


def add_two_col_table(doc, rows, headers=("Sales moment", "Language to use")):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(2.08)
    table.columns[1].width = Inches(4.34)
    hdr = table.rows[0].cells
    hdr[0].width = Inches(2.08)
    hdr[1].width = Inches(4.34)
    hdr[0].text = headers[0]
    hdr[1].text = headers[1]
    for cell in hdr:
        set_cell_fill(cell, "102019")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run(run, size=8.5, color=WHITE, bold=True)
    set_repeat_table_header(table.rows[0])
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(2.08)
        cells[1].width = Inches(4.34)
        for cell in cells:
            set_cell_margins(cell, top=135, bottom=135)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cells[0], "E7DDC9")
        set_run(cells[0].paragraphs[0].add_run(left), size=9.3, color=FOREST, bold=True)
        set_run(cells[1].paragraphs[0].add_run(right), size=9.3, color=INK)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def build():
    doc = Document()
    style_document(doc)
    props = doc.core_properties
    props.title = "Cinematic Flight Sales Playbook"
    props.subject = "Editable sales pitch and conversation guide"
    props.author = "Cinematic Flight"
    props.keywords = "Cinematic Flight, property websites, hospitality, sales pitch"

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(18)
    logo_shape = p.add_run().add_picture(str(LOGO), width=Inches(3.25))
    logo_shape._inline.docPr.set("descr", "Cinematic Flight threshold-and-path logo")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run("SALES PLAYBOOK"), size=9, color=BRASS_DARK, bold=True)
    p.runs[0].font.letter_spacing = Pt(1.4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(9)
    set_run(
        p.add_run(
            "Scroll-controlled cinematic property websites\n"
            "built from the photographs a property already has."
        ),
        font="Bodoni Moda",
        size=22,
        color=FOREST,
    )
    p.paragraph_format.line_spacing = 1.04

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(17)
    set_run(p.add_run("A living sales pitch for visually distinctive hospitality properties"), size=12, color=BRASS_DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    hero_shape = p.add_run().add_picture(str(HERO), width=Inches(6.18))
    hero_shape._inline.docPr.set("descr", "Good Food Farm surrounded by tropical greenery")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    set_run(p.add_run("You provide the place. We create the website and the cinematic experience."), font="Bodoni Moda", size=15, color=FOREST)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Working demonstration: Good Food Farm  •  Version 1.0  •  28 August 2026"), size=8.4, color=SAGE)

    page_break(doc)

    # Page 2: core pitch
    add_kicker(doc, "01  The core pitch")
    add_h1(doc, "Make the place felt before the guest arrives.")
    add_body(doc, "Cinematic Flight is a website design and build service for resorts, retreats, villas, farm stays, hospitality venues, and destination properties whose current website does not communicate the quality of the place.")
    add_quote(doc, "You provide the place. We create the website and the cinematic experience.", "Primary message")

    add_h2(doc, "The 30-second pitch")
    add_body(doc, "Most property websites reduce a distinctive place to the same grid of rooms, rates, and thumbnails. Cinematic Flight takes the photographs a property already has and turns them into a responsive, scroll-controlled journey inside a complete property website. Instead of asking a visitor to imagine the place from a gallery, the website lets them move through it—then connects that experience to the property’s enquiry or booking route.")

    add_h2(doc, "The one-line version")
    add_body(doc, "We design property websites you can enter, using real photographs of the real place.", italic=True)

    add_h2(doc, "The first next step")
    add_body(doc, "Start with a free reading of the property’s existing photographs. We assess whether the material can carry a convincing journey, name any missing views, and explain what the website could become—before anything is designed or agreed.")
    add_note(doc, "Sales discipline", "Do not lead with technology, AI, animation, or production method. Lead with the gap the owner recognises: the place is better than the website. Then let the working Good Food Farm demonstration make the difference visible.")

    # Page 3: narrative
    page_break(doc)
    add_kicker(doc, "02  The sales narrative")
    add_h1(doc, "A conversation that moves from recognition to action.")
    add_number(doc, "Name the gap", "“Your place is better than your website.” Most property sites present a place like a catalogue. The photographs may be strong; the template flattens them.")
    add_number(doc, "Make the consequence human", "A person choosing where to stay is trying to feel the place before committing. A thumbnail gallery asks them to do the emotional work themselves.")
    add_number(doc, "Show the same material two ways", "Use the Good Food Farm demonstration: the same verified photographs appear first as a familiar gallery and then as an ordered cinematic journey. Nothing is presented as a fabricated competitor before-and-after.")
    add_number(doc, "Explain the service", "We select the real photographs that carry the character and flow of the property, create the cinematic view, and design the complete website around it. The result is responsive on desktop and mobile.")
    add_number(doc, "Remove the first objection", "New photography is not automatically required. Existing website, brochure, and archive images are reviewed first. If an essential view is missing, it is named clearly.")
    add_number(doc, "Offer the low-friction next step", "Invite the owner to a free reading of the photographs. The aim is not to force a project decision; it is to establish whether the material and property are a strong fit.")

    add_h2(doc, "Three-part explanation of the offer")
    add_bullet(doc, "Your property — we identify photographs that establish the place and its flow.", "Your property")
    add_bullet(doc, "The cinematic view — we connect those photographs into a guided, scroll-controlled journey.", "The cinematic view")
    add_bullet(doc, "The complete website — we design the commercial experience around the journey and connect the enquiry or booking route.", "The complete website")
    add_note(doc, "Delivery choice", "The cinematic view may be added to a suitable current website, or it may become the signature experience inside a complete new property website. State this once, after the three-part explanation.")

    # Page 4: meeting flow
    page_break(doc)
    add_kicker(doc, "03  The discovery conversation")
    add_h1(doc, "Qualify the place before selling the build.")
    add_h2(doc, "Opening questions")
    for item in [
        "What do guests understand immediately when they arrive at the property that your current website fails to communicate?",
        "Which views, spaces, or transitions make people say, “I didn’t realise it was like this”?",
        "Where do your strongest photographs live today—website, brochure, booking platforms, social media, or an archive?",
        "What action should a visitor take after they feel the place: enquire, check availability, book, request a package, or contact a host?",
        "Are you improving the current website or open to a complete rebuild?",
    ]:
        add_bullet(doc, item)

    add_h2(doc, "Signals of a strong fit")
    add_two_col_table(doc, [
        ("Distinctive physical experience", "The place has a strong arrival, landscape, sequence of spaces, or atmosphere that static thumbnails undersell."),
        ("Usable real photography", "The owner already has several coherent views with enough quality and continuity to assess."),
        ("Visible website gap", "The current site does not express the quality, character, or flow the owner knows guests experience in person."),
        ("Clear visitor action", "The journey can lead naturally to an enquiry, booking route, availability check, or another approved conversion path."),
        ("Decision-maker access", "The conversation includes the owner or stakeholder who can decide on brand, website, and commercial direction."),
    ], headers=("Fit signal", "What to listen for"))

    add_h2(doc, "Suggested 20-minute demo flow")
    add_number(doc, "3 minutes", "Ask what the current website fails to convey.", style="List Number 2")
    add_number(doc, "5 minutes", "Show the immersive Good Food Farm opening without over-explaining it.", style="List Number 2")
    add_number(doc, "4 minutes", "Show the same photographs as a gallery and as a flight.", style="List Number 2")
    add_number(doc, "4 minutes", "Explain property → cinematic view → complete website.", style="List Number 2")
    add_number(doc, "4 minutes", "Invite the free photographic reading and agree the next exchange.", style="List Number 2")

    # Page 5: objections
    page_break(doc)
    add_kicker(doc, "04  Objections and responses")
    add_h1(doc, "Answer plainly. Keep the experience premium and the claims grounded.")
    add_two_col_table(doc, [
        ("“Do we need a new photo shoot?”", "Not necessarily. We begin with the photographs you already have and assess their suitability before anything is built. If a key view is missing, we name it clearly."),
        ("“Is this only an animation?”", "No. Cinematic Flight is a website design and build service. The cinematic property view is the signature experience inside the website, supported by the information and action a guest needs."),
        ("“Can it work with our current site?”", "Possibly. The cinematic view can be considered for a suitable existing website, or it can anchor a complete rebuild. The right route depends on the current site and commercial goal."),
        ("“Will it work on phones?”", "The working prototype uses separate desktop and phone-weight media and responsive layouts. Production scope still includes real-device verification for the final property experience."),
        ("“Will it improve bookings?”", "The purpose is to communicate the place more powerfully and connect that experience to the property’s chosen action. We do not promise a booking uplift without measured client results."),
        ("“Is this AI-generated imagery?”", "The experience is anchored in verified photographs of the real property. The site does not use synthetic imagery to represent the place."),
        ("“How much does it cost?”", "Pricing is defined after the photographic reading and a clear scope: add to an existing site or build a complete new property website. Do not quote an unapproved range."),
    ])

    add_note(doc, "Prototype truth", "The Good Food Farm example is a working demonstration. The local enquiry form currently demonstrates validation and success states only; it is not connected to an external endpoint. Do not describe it as a live lead-capture system until that connection is approved and implemented.")

    # Page 6: approaching a lead
    page_break(doc)
    add_kicker(doc, "05  Approaching a lead")
    add_h1(doc, "Start with the place, not the sale.")
    add_body(doc, "The strongest opening shows that you noticed something specific about the property. The lead should feel selected because the place has character—not added to a generic web-design list.")

    add_h2(doc, "The spoken spiel")
    add_quote(doc, "Hi [NAME]. I came across [PROPERTY], and [SPECIFIC VIEW, ARRIVAL, LANDSCAPE, OR ATMOSPHERE] stood out. The place feels distinctive, but a normal gallery can only show it one frame at a time. I design scroll-controlled cinematic property websites using photographs you already have, so visitors can move through the place before arriving. I’d be happy to give you a free reading of the images already online and show what they might do. Would that be useful?", "45-second conversation opener")

    add_h2(doc, "A simple five-part approach")
    add_number(doc, "Notice", "Choose one truthful, property-specific detail from the website or approved public material.", style="List Number 3")
    add_number(doc, "Recognise", "Name what makes the place feel distinctive without criticising the owner or insulting the current site.", style="List Number 3")
    add_number(doc, "Explain", "Say plainly that Cinematic Flight designs and builds property websites with a scroll-controlled cinematic experience made from existing photographs.", style="List Number 3")
    add_number(doc, "Demonstrate", "Offer the Good Food Farm example only after the lead understands the business offer.", style="List Number 3")
    add_number(doc, "Invite", "Ask permission to provide the free photographic reading. Use one small, clear next action.", style="List Number 3")

    add_h2(doc, "Tone guardrails")
    add_bullet(doc, "Personalise the first two sentences; never send an observation you cannot support.")
    add_bullet(doc, "Say “the website may not communicate the full feeling of the place,” not “your website is bad.”")
    add_bullet(doc, "Avoid AI, animation, drone, and production jargon in the opening message.")
    add_bullet(doc, "Do not promise more bookings, a delivery date, or a price before the offer is defined.")
    add_bullet(doc, "End with one question that is easy to answer.")

    # Pages 7-8: message library
    page_break(doc)
    add_kicker(doc, "06  Lead message library")
    add_h1(doc, "Messages you can copy, personalise, and send.")
    add_note(doc, "Before sending", "Replace every bracketed field, keep one specific observation, and choose only one call to action. If the message could be sent unchanged to another property, it is not ready.")

    add_h2(doc, "Warm contact")
    add_body(doc, "Hi [NAME]—I’m building Cinematic Flight, a website design service for distinctive properties. It turns photographs a property already has into a scroll-controlled journey inside the website. I thought of [PROPERTY] because [SPECIFIC REASON]. Would you be open to a quick look at the working example?", italic=True)

    add_h2(doc, "Referral introduction")
    add_body(doc, "Hi [NAME], [REFERRER] suggested I contact you. I’ve been looking at [PROPERTY], especially [SPECIFIC DETAIL]. I create cinematic property websites from existing photographs, helping the website express more of what the place feels like in person. I can start with a free reading of the photographs already online. Would it be useful if I sent a short example?", italic=True)

    add_h2(doc, "Cold email")
    add_body(doc, "Subject: An idea for how [PROPERTY] could feel online", bold_lead="Subject: ")
    add_body(doc, "Hi [NAME], I came across [PROPERTY] and noticed [SPECIFIC, TRUTHFUL OBSERVATION]. The place appears to have a stronger sense of arrival and atmosphere than a conventional photo gallery can communicate. I design and build scroll-controlled cinematic property websites using photographs the property already has. I’d be happy to give you a free reading of the images already online and tell you whether they could carry this kind of experience. May I send you the working Good Food Farm example?", italic=True)

    add_h2(doc, "Instagram or Facebook message")
    add_body(doc, "Hi—[SPECIFIC DETAIL] at [PROPERTY] caught my attention. I build cinematic property websites that turn existing photographs into a guided, scroll-controlled journey. I think your material may be a strong fit. Would you like me to send a short working example and a free reading of the photographs already online?", italic=True)

    add_h2(doc, "LinkedIn message")
    add_body(doc, "Hi [NAME], I’m reaching out because [PROPERTY] has [SPECIFIC QUALITY], and the current online presentation may not communicate the full feeling of the place. Cinematic Flight designs property websites around a scroll-controlled journey made from existing photographs. If useful, I can send a concise example and assess whether the property’s current images could carry it.", italic=True)

    page_break(doc)
    add_kicker(doc, "06  Message library continued")
    add_h1(doc, "Keep the next step small and specific.")

    add_h2(doc, "When the lead shows interest")
    add_body(doc, "Great—thank you. The next step is for me to read [PROPERTY]’s existing photographs, see whether they establish a convincing sequence, and identify any important missing view. Please send the website or approved photo link you would like me to review.", after=5, italic=True)

    add_h2(doc, "Invitation to a short demo")
    add_body(doc, "I can show you the Good Food Farm demonstration in 15–20 minutes: the experience, the same photographs as a gallery, and how the journey connects to a complete website. Would [OPTION A] or [OPTION B] suit you?", after=5, italic=True)

    add_h2(doc, "First follow-up")
    add_body(doc, "Hi [NAME], just bringing this back to the top of your messages. I’m offering a free reading of [PROPERTY]’s existing photographs—whether they could support a cinematic website experience and what may be missing. Would you like me to take a look?", after=5, italic=True)

    add_h2(doc, "Second follow-up with value")
    add_body(doc, "Hi [NAME], one reason [PROPERTY] could suit this approach is [OBSERVATION]. That [VIEW / TRANSITION / ATMOSPHERE] could become part of the website journey instead of remaining one gallery thumbnail. I’m happy to send the working example if useful.", after=5, italic=True)

    add_h2(doc, "After a demo")
    add_body(doc, "Thank you for exploring Cinematic Flight. The useful next step is a short reading of [PROPERTY]’s existing photographs. I’ll assess whether they can form a convincing journey and name any missing views before we discuss scope. There is no commitment attached.", after=5, italic=True)

    add_h2(doc, "Graceful close-out")
    add_body(doc, "Hi [NAME], I’ll close the loop here so I don’t keep filling your inbox. If communicating the feeling of [PROPERTY] becomes a priority later, I’d be glad to provide the photographic reading or working example. Wishing you and the team well.", after=5, italic=True)

    add_h2(doc, "Useful calls to action")
    add_bullet(doc, "May I send the working example?")
    add_bullet(doc, "Would you like a free reading of the photographs already online?")
    add_bullet(doc, "Would a 15–20 minute walkthrough be useful?")
    add_bullet(doc, "Which website or approved photo folder should I review?")

    # Page 9: living fields
    page_break(doc)
    add_kicker(doc, "07  Living offer worksheet")
    add_h1(doc, "Update these fields as the commercial offer becomes real.")
    add_body(doc, "The brass fields are intentionally editable. Keep the core positioning stable; update only when there is approved evidence, a defined delivery process, or a deliberate commercial decision.")

    add_h2(doc, "Current offer")
    add_editable(doc, "Primary market", "Define the first geographic or hospitality niche")
    add_editable(doc, "Named service", "Confirm Cinematic Property Website, Cinematic Flight, or package name")
    add_editable(doc, "Delivery options", "Define add-on criteria and complete-rebuild scope")
    add_editable(doc, "Pricing", "Add only after scope and commercial approval")
    add_editable(doc, "Lead time", "Add only after production capacity is validated")
    add_editable(doc, "Enquiry route", "Insert approved email, form endpoint, or booking link")

    add_h2(doc, "Proof ledger")
    add_editable(doc, "Working demonstrations", "Good Food Farm; add only verified work")
    add_editable(doc, "Client outcomes", "Insert measured results with source and date")
    add_editable(doc, "Testimonials", "Insert approved quotation and permission status")
    add_editable(doc, "Device verification", "Record production browser and real-device checks")

    add_h2(doc, "Never drift from these truths")
    add_bullet(doc, "Describe Cinematic Flight first as a website design and build service.")
    add_bullet(doc, "Use real property photographs as the visible proof; do not introduce synthetic property imagery.")
    add_bullet(doc, "Do not invent performance results, testimonials, client outcomes, pricing, availability, or production promises.")
    add_bullet(doc, "Keep the free photographic reading as the low-friction first step until the offer changes deliberately.")
    add_bullet(doc, "Mark prototype features as prototypes until the live service is implemented and verified.")

    add_note(doc, "Revision rule", "When a claim changes, record the evidence, source, owner, and approval date beside it before using the claim in sales material.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
